# -*- coding: utf-8 -*-
"""书面转正答辩报告（可归档、可口述）"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.4)
sec.left_margin = Cm(2.6)
sec.right_margin = Cm(2.6)


def set_font(run, size=12, bold=False, color=None, name="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    if color:
        run.font.color.rgb = color
    r = run._element.get_or_add_rPr()
    r.rFonts.set(qn("w:eastAsia"), name)


def add_p(text, size=12, bold=False, center=False, space_after=8, first_line=None, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.35
    if first_line is not None:
        p.paragraph_format.first_line_indent = Cm(first_line)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size, bold, color)
    return p


navy = RGBColor(0x1B, 0x36, 0x5D)

add_p("销售与技术服务部", 12, True, True, 4, color=navy)
add_p("CIS 销售经理试用期转正答辩报告", 18, True, True, 8, color=navy)
add_p("负责客户：华勤（手机 ODM）  ·  销售公司 CIS 芯片  ·  2026 年 8 月", 11, False, True, 16, color=RGBColor(0x5A, 0x6A, 0x7A))

add_p("汇报说明", 14, True, False, 8, color=navy)
add_p(
    "本报告对应市场营销例会「销售与技术服务部」周报口径：先讲本周期（试用期）实际做了什么，再评价是否胜任，最后给出可被周报检查的下阶段计划。岗位职责是负责华勤手机业务，销售公司 CIS 芯片。华勤是手机 ODM，品牌方管规格与 AVL，华勤管落地节奏与第二源建议，模组厂管工程验证。试用期的合理产出是「可推进的项目」——窗口、漏斗、送样与问题闭环，而不是尚未发生的定点金额。",
    first_line=0.74,
)

add_p("一、试用期工作总结", 14, True, False, 8, color=navy)
add_p(
    "试用期我按「先结构、后项目、再机制」推进，避免一上来只追一颗料。工作可以收成四条线。",
    first_line=0.74,
)
add_p("（一）客户经营：把华勤从「一个名字」变成可经营的账号", 12, True, space_after=6)
add_p(
    "完成华勤手机 CIS 相关决策链梳理，覆盖项目经理、摄像头硬件、影像/调试、采购、质量与计划。针对不同角色准备不同材料：项目要节点和风险，硬件要规格/封装/替代可行性，影像要验证结论，采购要价格与供给，质量要导入风险。建立采购与项目双线窗口，并安排备份对接，沟通节奏按周或双周维持，拜访形成纪要。同时明确约束：华勤可以推荐和推进验证，但品牌 AVL 不打开，项目到不了定点。试用期把 ODM 内部盘清楚，是后续做品牌侧的前提，而不是回避品牌。",
    first_line=0.74,
)
add_p("（二）项目导入：用漏斗管理，而不是「跟一颗料」", 12, True, space_after=6)
add_p(
    "输出华勤在研机型摄像头规格与竞品卡位表，看清格科、思特威、豪威等成熟料在哪些像素档占位、开案窗口在哪。主攻卡位放在手机前置及功能摄等 2M 规格档，对标成熟 1/5\" 2M 料做第二源，话术强调供给弹性、导入服务和风险对冲，不硬刚纸面参数。组织目标规格送样，协同 FAE 完成陪访与问题闭环，资料包（规格对照、setting、FAQ）随项目走。项目按 A/B/C 分级进入周报，红黄绿状态对内对齐。试用期实际跑通的是情报—匹配—送样，目标项目进入或准备进入验证；定点与量产纳入漏斗管理，不提前报喜。品牌 AVL 路径已识别为红灯，未打开窗口。",
    first_line=0.74,
)
add_p("（三）内部协同：销售对节点负责，技术服务对结论负责", 12, True, space_after=6)
add_p(
    "与 FAE 约定：拜访前出问题清单，现场销售控议程、FAE 答技术，当天纪要写清问题、责任人和截止日，技术问题 48 小时书面回复，卡住节点当周升级。对内对外同一套事实：客户要的是能否赶上这台机，内部要的是问题能否关闭。导入风险（封装、CRA、OTP、上电时序等）提前讲，不留到验证后期。竞品动态进周报，供产品与计划决策。不把送样次数写成业绩。",
    first_line=0.74,
)
add_p("（四）专业内功", 12, True, space_after=6)
add_p(
    "补齐手机向 CIS 的规格语言、竞品格局和华勤 EVT/DVT/MP 节点认知，掌握送样、模组验证和 OTP 等导入节奏，能够独立完成常规拜访与纪要。复杂 IQ 与像素细节仍与 FAE 协同，不抢技术结论。",
    first_line=0.74,
)
add_p("试用期可核验的结果包括：决策链地图、主对接与备份窗口、规格与竞品卡位表、送样与 FAE 协同机制、A/B/C 漏斗与周报、华勤账号作战手册初稿。未完成项主要是品牌 AVL 窗口未打开、尚无量产定点、华勤专用商务口径未锁死、售前问答仍部分依赖 FAE。", first_line=0.74)

add_p("二、试用期自我评价", 14, True, False, 8, color=navy)
add_p(
    "综合评价：匹配 CIS 销售经理（华勤手机）这个岗位。能够在多层决策的 ODM 账号上，把芯片销售做成可管理的项目，而不是个人关系或一次性送样。这是转正的基本盘。经营结果（定点、收入）需要在转正后的验证和商务阶段兑现，试用期不夸大。",
    first_line=0.74,
)
add_p("分维度自评如下。客户洞察强：能把品牌—华勤—模组三层结构讲清楚，知道谁决策、谁影响、哪里会卡住。沟通协同强：对内对外同一套事实，与 FAE 的分工可执行。抗压担当强：长周期、多角色，问题当周暴露，不藏到月底。目标管理中偏上：漏斗和周报已建立，A 类项目的周节点还要更细。结果导向中偏上：过程结果扎实，经营结果尚未发生。专业能力中等：规格与导入流程已入门，OTP/IQ 细节还不能独立覆盖。", first_line=0.74)
add_p("适合留下的理由：不把 ODM 当渠道批发；执行闭环清楚（纪要、责任人、升级）；对数字诚实，不把漏斗当业绩；愿意补专业，不与 FAE 抢结论。必须补齐的短板：品牌侧渗透几乎为零；独立商务谈判经验不足；技术问答还不能独立覆盖约 70%；A 类项目周节点仍偏粗。这些短板都写入 90 天计划，而不是用态度带过。", first_line=0.74)
add_p("对「为什么试用期还没有定点」的回答：手机 CIS 从情报到定点通常跨 EVT/DVT，周期长于试用期；竞品已在品牌短名单，第二源验证需要模组与 IQ 结论。因此试用期 KPI 应看窗口质量、漏斗是否真实、验证是否在推进、风险是否提前暴露。若用尚未发生的收入评价试用期，会倒逼报喜或空许诺，这与周报「红灯当周升级」的要求相反。", first_line=0.74)

add_p("三、未来规划与展望", 14, True, False, 8, color=navy)
add_p("（一）转正后 90 天：四件必须发生的事", 12, True, space_after=6)
add_p(
    "第一，验证出结论。至少一个目标项目完成模组或 IQ 验证，形成书面结论：过、有条件过、或不过及原因。第二，打开品牌线索。画出品牌—华勤—模组三方图，拿到至少一个品牌侧或品牌指定模组的有效窗口线索。第三，锁定商务口径。与产品、计划形成华勤专用报价、交期、供给原则，现场不再口头承诺。第四，机制固化。华勤项目月评覆盖漏斗、风险、FAE 资源和下一步动作，账号经营公司看得见，并继续用周报红黄绿接受检查。",
    first_line=0.74,
)
add_p("（二）未来 12 个月：把华勤做成样板账号", 12, True, space_after=6)
add_p(
    "样板的意思是方法可复制，而不只是「我跟过一个客户」。项目上，争取 1 个项目进入试产或定点路径，2–3 个项目稳定在验证池，杜绝只有拜访没有下一节点的空账号。关系上，从单一窗口升级到项目组级：采购、硬件、影像、质量都有可对话的人，且有备份。产品上，以 2M 等可切入规格做第二源破口，再向更高规格在研机型要验证机会，不一次铺太多料号。组织上，沉淀《华勤 CIS 导入包》（规格对照、setting、OTP/封装注意事项、FAQ、竞品卡位），可交接。若公司有规划，将打法复制到同类手机 ODM，使销售经理从管一个客户变成管一类战场。能力上，独立完成约 70% 售前，复杂 IQ 仍协同 FAE，并能主持华勤月度项目会。",
    first_line=0.74,
)
add_p("（三）承诺与支持诉求", 12, True, space_after=6)
add_p(
    "我承诺：华勤项目状态进入周报，红灯当周升级；对外口径与对内事实一致；不把送样数当业绩，只认下一节点；90 天四件事按月复盘，未完成说明原因和下一步；与 FAE 按约定分工，不甩锅、不抢功。请公司支持：产品明确可主推华勤的规格与 roadmap；FAE 在验证期保证档期、问题 48 小时响应；商务给出能进短名单的第二源价格与供给策略；必要时协助品牌侧协同，避免只在 ODM 空转；质量与交付的导入风险清单与我同步对外讲。支持事项都对应 90 天目标，不是要资源的空话。",
    first_line=0.74,
)

add_p("结束语", 14, True, False, 8, color=navy)
add_p(
    "试用期把华勤的结构和漏斗立起来了。申请转正后，用验证结论和定点路径，把这个账号做成公司手机 CIS 的样板，并按销售与技术服务部周报机制接受检查。请批评指正。",
    first_line=0.74,
)

add_p("汇报人：________________    日期：2026 年 ____ 月 ____ 日", 12, False, True, 20)

out = "/workspace/CIS销售经理-华勤手机业务-转正答辩报告.docx"
doc.save(out)
print("saved", out)
